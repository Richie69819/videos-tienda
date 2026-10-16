# -*- coding: utf-8 -*-
"""Genera el guion de una clase en Word, con formato para leer mientras se graba."""
import sys, io, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
NARANJA = RGBColor(0xC2, 0x56, 0x0F)
GRIS = RGBColor(0x60, 0x60, 0x60)
NEGRO = RGBColor(0x1A, 0x1A, 0x1A)


def sombrear(par, hexcolor):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    pPr.append(shd)


def borde_izq(par, hexcolor, size=18):
    pPr = par._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), str(size))
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), hexcolor)
    pbdr.append(left)
    pPr.append(pbdr)


def base(doc):
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    st.font.color.rgb = NEGRO
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for s in doc.sections:
        s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)


def pie(doc, texto):
    p = doc.sections[0].footer.paragraphs[0]
    p.text = texto
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.size = Pt(8); r.font.color.rgb = GRIS


def titulo(doc, texto, sub=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(texto)
    r.font.size = Pt(26); r.bold = True; r.font.color.rgb = AZUL
    if sub:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(14)
        r2 = p2.add_run(sub)
        r2.font.size = Pt(12); r2.font.color.rgb = GRIS


def parte(doc, numero, texto, dur):
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    sombrear(p, '1F3A5F')
    r = p.add_run('  PARTE %s  ' % numero)
    r.font.size = Pt(13); r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r2 = p.add_run('   %s' % texto)
    r2.font.size = Pt(13); r2.bold = True
    r2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if dur:
        p3 = doc.add_paragraph()
        p3.paragraph_format.space_after = Pt(8)
        p3.paragraph_format.keep_with_next = True
        r3 = p3.add_run('Dura como %s' % dur)
        r3.font.size = Pt(9); r3.italic = True; r3.font.color.rgb = GRIS


def rotulo(doc, texto, color):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(texto)
    r.font.size = Pt(10); r.bold = True; r.font.color.rgb = color
    return p


def haces(doc, pasos):
    rotulo(doc, 'LO QUE HACES EN PANTALLA', NARANJA)
    for i, t in enumerate(pasos):
        p = doc.add_paragraph(style='List Number' if False else None)
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run('%d.  ' % (i + 1))
        r.bold = True; r.font.color.rgb = NARANJA; r.font.size = Pt(10.5)
        r2 = p.add_run(t)
        r2.font.size = Pt(10.5)


def dices(doc, lineas):
    rotulo(doc, 'LO QUE DICES', AZUL)
    for t in lineas:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(7)
        borde_izq(p, '1F3A5F')
        r = p.add_run(t)
        r.font.size = Pt(12.5)


def nota(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    sombrear(p, 'F2F2F2')
    r = p.add_run('NOTA PARA TI (no se dice):  ')
    r.font.size = Pt(9.5); r.bold = True; r.font.color.rgb = GRIS
    r2 = p.add_run(texto)
    r2.font.size = Pt(9.5); r2.italic = True; r2.font.color.rgb = GRIS


def h2(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(texto)
    r.font.size = Pt(14); r.bold = True; r.font.color.rgb = AZUL


def parrafo(doc, texto, size=11):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.font.size = Pt(size)
    return p


def casillas(doc, items):
    for t in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run('☐  ')
        r.font.size = Pt(13)
        r2 = p.add_run(t)
        r2.font.size = Pt(11)


def tabla_atajos(doc, filas):
    t = doc.add_table(rows=1, cols=3)
    t.style = 'Light Grid Accent 1'
    hdr = t.rows[0].cells
    for c, txt in zip(hdr, ['Tecla', 'Qué hace', 'Cómo se recuerda']):
        c.text = ''
        r = c.paragraphs[0].add_run(txt)
        r.bold = True; r.font.size = Pt(10)
    for tecla, que, como in filas:
        cells = t.add_row().cells
        r = cells[0].paragraphs[0].add_run(tecla); r.bold = True; r.font.size = Pt(11)
        r = cells[1].paragraphs[0].add_run(que); r.font.size = Pt(10.5)
        r = cells[2].paragraphs[0].add_run(como); r.font.size = Pt(10); r.font.color.rgb = GRIS
    return t
