# -*- coding: utf-8 -*-
"""Helpers del curso corto."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from guion_docx import (base, pie, titulo, h2, parrafo, casillas, sombrear,
                        GRIS, AZUL, NARANJA)

SALIDA = r"C:\Users\RICHO\Desktop\Curso Photoshop\guiones"
ROJO = RGBColor(0xB0, 0x22, 0x22)

DISCIPLINA = ("Esta clase son {dur}. Si te pasas de {tope}, no es que te falte decir algo: "
              "es que te alargaste en una parte. Corta y sigue; lo que se te qued\u00f3 va a "
              "los comentarios.")


def nuevo(tit, sub, piecito):
    doc = Document()
    base(doc)
    pie(doc, piecito)
    titulo(doc, tit, sub)
    for etiqueta, resto, color in [
        ("LO QUE HACES EN PANTALLA", " \u2014 los clics, en orden.", NARANJA),
        ("LO QUE DICES", " \u2014 el texto hablado. El orden no se cambia.", AZUL),
        ("NOTA PARA TI", " \u2014 nunca se dice en voz alta.", GRIS),
    ]:
        pp = doc.add_paragraph()
        pp.paragraph_format.left_indent = Cm(0.6)
        pp.paragraph_format.space_after = Pt(3)
        r = pp.add_run(etiqueta); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = color
        r2 = pp.add_run(resto); r2.font.size = Pt(10)
    return doc


def regla(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    sombrear(p, 'FFF4E5')
    r = p.add_run("  " + texto + "  ")
    r.font.size = Pt(11); r.bold = True
    return p


def reloj(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("RELOJ:  ")
    r.font.size = Pt(10); r.bold = True; r.font.color.rgb = ROJO
    r2 = p.add_run(texto)
    r2.font.size = Pt(10); r2.font.color.rgb = ROJO


def antes(doc, items):
    h2(doc, "Antes de darle Rec")
    casillas(doc, items)


def cierre_notas(doc, titulo_yt, entrega, miniatura, shorts):
    doc.add_page_break()
    h2(doc, "Notas de producci\u00f3n")
    parrafo(doc, "Nada de esto se dice en c\u00e1mara.", 10)
    parrafo(doc, "T\u00edtulo: " + titulo_yt)
    parrafo(doc, "Entrega: " + entrega)
    parrafo(doc, "Miniatura: " + miniatura)
    h2(doc, "Shorts de esta grabaci\u00f3n")
    for t in shorts:
        p = doc.add_paragraph(t, style='List Bullet')
        p.paragraph_format.left_indent = Cm(0.8)


def guardar(doc, nombre):
    ruta = os.path.join(SALIDA, nombre)
    doc.save(ruta)
    print("DOCX:", ruta)
